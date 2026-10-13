"""One-off export of the frozen NQ 3216 Diagnosis inputs from the annotation workbook.

Provenance only: this is the single script that reads the local, gitignored
workbook data/annotations/NQ_C1_3216_initial_labels.xlsx. Its output,
data/annotations/normalized/nq_3216_diagnosis_inputs.jsonl, is the committed
source artifact used by scripts/build_nq_2698_diagnosis_trusted.py; the formal
v2 pipeline never needs the workbook.

Only the runtime columns sample_id, question, c1_prediction and top5_evidence
of the main sheet are read. gold_answer, em_reference and the label columns
are never read. Values are exported verbatim in workbook order:
- c1_prediction keeps the raw cell string (e.g. the literal model output
  "None"); empty-prediction normalisation happens in the dataset builder.
- top5_evidence is split into {rank, passage_id, score, text}. Every workbook
  score has exactly four decimals, so the original block is recovered as
  "\\n\\n".join(f"[{rank}] {passage_id} | score={score:.4f}\\n{text}"); the
  export checks this round trip for every record before writing.

If the output file already exists it is never overwritten: the exported
records are compared line by line with it and the script fails on any
difference.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

WORKBOOK = ROOT / "data" / "annotations" / "NQ_C1_3216_initial_labels.xlsx"
MAIN_SHEET = "NQ C1 Annotation"
RUNTIME_COLUMNS = ("sample_id", "question", "c1_prediction", "top5_evidence")
OUTPUT_PATH = ROOT / "data" / "annotations" / "normalized" / "nq_3216_diagnosis_inputs.jsonl"
EXPECTED_ALL = 3216
EVIDENCE_HEADER = re.compile(r"^\[(\d+)\] (\S+) \| score=(\S+)\n", re.M)


def read_runtime_rows() -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        rows = workbook[MAIN_SHEET].iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
        positions = [header.index(column) for column in RUNTIME_COLUMNS]
        return [
            {column: row[position] for column, position in zip(RUNTIME_COLUMNS, positions)}
            for row in rows
            if any(cell not in (None, "") for cell in row)
        ]
    finally:
        workbook.close()


def parse_evidence(block: str) -> list[dict]:
    headers = list(EVIDENCE_HEADER.finditer(block))
    evidence = []
    for i, match in enumerate(headers):
        end = headers[i + 1].start() if i + 1 < len(headers) else len(block)
        evidence.append({
            "rank": int(match.group(1)),
            "passage_id": match.group(2),
            "score": float(match.group(3)),
            "text": block[match.end():end].strip(),
        })
    if [e["rank"] for e in evidence] != [1, 2, 3, 4, 5]:
        raise ValueError(f"Expected evidence ranks 1..5, got {[e['rank'] for e in evidence]}")
    return evidence


def render_evidence(evidence: list[dict]) -> str:
    return "\n\n".join(f"[{e['rank']}] {e['passage_id']} | score={e['score']:.4f}\n{e['text']}" for e in evidence)


def export_row(row: dict) -> dict:
    for column in RUNTIME_COLUMNS:
        if not isinstance(row[column], str):
            raise ValueError(f"{row['sample_id']}: {column} is not a string cell: {row[column]!r}")
    evidence = parse_evidence(row["top5_evidence"])
    if render_evidence(evidence) != row["top5_evidence"]:
        raise ValueError(f"{row['sample_id']}: top5_evidence does not round-trip losslessly")
    return {
        "sample_id": row["sample_id"],
        "question": row["question"],
        "c1_prediction": row["c1_prediction"],
        "top5_evidence": evidence,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    args = parser.parse_args()

    records = [export_row(row) for row in read_runtime_rows()]
    ids = [record["sample_id"] for record in records]
    if len(records) != EXPECTED_ALL or len(set(ids)) != len(ids):
        raise ValueError(f"Expected {EXPECTED_ALL} unique sample IDs, got {len(records)} ({len(set(ids))} unique)")
    lines = [json.dumps(record, ensure_ascii=False) for record in records]

    if args.output.exists():
        existing = args.output.read_text(encoding="utf-8").splitlines()
        diffs = [i for i, (a, b) in enumerate(zip(lines, existing)) if a != b]
        if len(existing) != len(lines) or diffs:
            raise SystemExit(f"Exported records differ from {args.output}: lines {len(lines)} vs {len(existing)}, "
                             f"first differing record index {diffs[:5]}")
        print(f"Verified: {len(lines)} exported records are identical to {args.output}")
    else:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
        print(f"Wrote {len(lines)} records to {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
